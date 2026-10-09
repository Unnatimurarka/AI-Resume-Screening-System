import docx

from conftest import make_resume

from screener.ingest import read_resume
from screener.extract import parse_resume


def test_basic_fields():
    r = make_resume("ASHA RAO\nasha.rao@example.com | +91 99999 99999\ngithub.com/asharao\n\nEXPERIENCE\n- Built a thing in Python.\n")
    assert r.candidate_name == "Asha Rao"
    assert r.email == "asha.rao@example.com"
    assert r.github_username == "asharao"
    assert "Python" in r.matched_skills


def test_profile_links_beat_repo_links_from_other_owners():
    r = make_resume("Jo Li\njo@example.com\n", links=["https://github.com/jo", "https://github.com/jo/proj", "https://github.com/someorg/lib"])
    assert r.github_username == "jo"


def test_no_github_gives_none():
    assert make_resume("Jo Li\njo@example.com\n").github_username is None


def test_skills_list_lines_are_tagged():
    r = make_resume("Jo Li\n\nPROJECTS\n\n- Built an agent in Python.\n\nTECHNICAL SKILLS\n\nLanguages: Python, Go, SQL, Rust\n")
    kinds = {u.text: u.kind for u in r.units}
    assert kinds["Built an agent in Python."] == "body"
    assert kinds["Languages: Python, Go, SQL, Rust"] == "skills"


def test_docx_is_supported(tmp_path):
    d = docx.Document()
    d.add_paragraph("Mira Shah")
    d.add_paragraph("mira@example.com")
    d.add_paragraph("Built a RAG pipeline in Python with LangChain and Qdrant for support tickets.")
    path = tmp_path / "mira.docx"
    d.save(path)
    r = parse_resume(read_resume(path))
    assert r.candidate_name == "Mira Shah" and r.email == "mira@example.com"
